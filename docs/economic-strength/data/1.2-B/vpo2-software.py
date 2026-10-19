#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""1.2-B: оснащённость вузов по форме ВПО-2 (2025).

Скачивает архив ВПО-2 за 2025 г., распаковывает региональные своды «*_ВСЕГО.xlsx»
и выходную таблицу tab19, суммирует по России:
  - подраздел 2.1 — персональные компьютеры и оборудование;
  - подраздел 2.2 — наличие специальных программных средств (число организаций);
  - подраздел 2.3 — максимальная скорость доступа к Интернету (число организаций);
  - численность студентов ВО (tab19) для нормировки.

Выход: vpo2-osnashchennost.csv (метрика; значение; единица; источник).
Запуск из корня репозитория:  python docs/economic-strength/data/1.2-B/vpo2-software.py
Зависимости: openpyxl.
"""
import csv
import os
import sys
import urllib.request
import zipfile

import openpyxl

sys.stdout.reconfigure(encoding="utf-8")

BASE = os.path.dirname(os.path.abspath(__file__))
RAW = os.path.abspath(os.path.join(BASE, "..", "..", "..", "..", "data", "economic-strength", "1.2-B", "raw"))
OUT_CSV = os.path.join(BASE, "vpo2-osnashchennost.csv")
URL = "https://minobrnauki.gov.ru/upload/2026/07/%D0%A1%D0%B2%D0%BE%D0%B4%D1%8B%20%D0%92%D0%9F%D0%9E-2%20%D0%B7%D0%B0%202025%20%D0%B3.zip"


def decode_name(name: str) -> str:
    if all(ord(ch) < 256 for ch in name):
        try:
            return name.encode("cp437").decode("cp866")
        except Exception:
            return name
    return name


def ensure_archive() -> str:
    os.makedirs(RAW, exist_ok=True)
    dst = os.path.join(RAW, "VPO2-2025.zip")
    if not os.path.exists(dst):
        print("скачивание:", URL)
        urllib.request.urlretrieve(URL, dst)
    print(f"архив: {os.path.getsize(dst):,} байт".replace(",", " "))
    return dst


def unpack(archive: str) -> str:
    out = os.path.join(RAW, "vpo2-2025")
    os.makedirs(out, exist_ok=True)
    with zipfile.ZipFile(archive) as z:
        for n in z.namelist():
            name = decode_name(n)
            if name.startswith("Выходные таблицы ВПО-2/") and name.endswith("tab19.xlsx"):
                dest = os.path.join(out, "tab19.xlsx")
            elif name.startswith("Своды ВПО-2/") and name.endswith(".xlsx") and "_ВСЕГО" in name:
                dest = os.path.join(out, os.path.basename(name))
            else:
                continue
            if not os.path.exists(dest):
                with z.open(n) as src, open(dest, "wb") as f:
                    f.write(src.read())
    files = sorted(f for f in os.listdir(out) if f.endswith("_ВСЕГО.xlsx"))
    print("региональных сводов:", len(files))
    return out


def as_num(v):
    if isinstance(v, (int, float)):
        return float(v)
    if isinstance(v, str):
        s = v.strip().replace(" ", "").replace(",", ".")
        try:
            return float(s)
        except ValueError:
            return None
    return None


def parse_region(path: str):
    """Возвращает счётчики Р2_2, Р2_3 по одной организации-региону."""
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    p22 = {}
    for row in wb["Р2_2"].iter_rows(min_row=7, max_row=18, values_only=True):
        label = str(row[0]).strip() if row[0] else ""
        present = as_num(row[2])
        avail = as_num(row[3])
        if label and present is not None:
            p22[label] = p22.get(label, 0.0) + present
            if avail is not None:
                p22[label + " (доступно обучающимся)"] = p22.get(label + " (доступно обучающимся)", 0.0) + avail

    p23 = {}
    for row in wb["Р2_3"].iter_rows(min_row=6, max_row=6, values_only=True):
        for i, v in enumerate(row[2:11], start=1):
            n = as_num(v)
            if n is not None:
                p23[i] = p23.get(i, 0.0) + n
    wb.close()
    return p22, p23


def main() -> int:
    archive = ensure_archive()
    folder = unpack(archive)

    files = sorted(f for f in os.listdir(folder) if f.endswith("_ВСЕГО.xlsx"))
    total22, total23 = {}, {}
    pc = {"total": 0.0, "edu": 0.0, "free": 0.0, "new": 0.0, "laptops": 0.0, "tablets": 0.0}
    for i, f in enumerate(files, start=1):
        p22, p23 = parse_region(os.path.join(folder, f))
        for k, v in p22.items():
            total22[k] = total22.get(k, 0.0) + v
        for k, v in p23.items():
            total23[k] = total23.get(k, 0.0) + v

        wb = openpyxl.load_workbook(os.path.join(folder, f), read_only=True, data_only=True)
        for row in wb["Р2_1"].iter_rows(min_row=8, max_row=22, values_only=True):
            label = str(row[0]).strip() if row[0] else ""
            if label.startswith("Персональные компьютеры – всего"):
                pc["total"] += as_num(row[2]) or 0
                pc["edu"] += as_num(row[3]) or 0
                pc["free"] += as_num(row[4]) or 0
            elif label.startswith("из них:") and "ноутбуки" in label:
                pc["laptops"] += as_num(row[2]) or 0
            elif label.startswith("планшетные компьютеры"):
                pc["tablets"] += as_num(row[2]) or 0
            elif label.startswith("поступившие в отчетном году"):
                pc["new"] += as_num(row[2]) or 0
        wb.close()
        if i % 20 == 0:
            print(f"  разобрано {i}/{len(files)}")

    students_vo = None
    tab19 = os.path.join(folder, "tab19.xlsx")
    if os.path.exists(tab19):
        wb = openpyxl.load_workbook(tab19, read_only=True, data_only=True)
        for row in wb.worksheets[0].iter_rows(values_only=True):
            if str(row[0]).strip() == "Российская Федерация":
                students_vo = sum((as_num(row[i]) or 0.0) for i in (2, 3, 4))
                break
        wb.close()

    rows = []

    def add(metric, value, unit, source):
        rows.append({"metric": metric, "value": "" if value is None else round(value, 1), "unit": unit, "source": source})

    add("Персональные компьютеры — всего", pc["total"], "шт.", "ВПО-2, Р2_1")
    add("ПК, используемые в учебных целях", pc["edu"], "шт.", "ВПО-2, Р2_1")
    add("ПК, доступные студентам в свободное время", pc["free"], "шт.", "ВПО-2, Р2_1")
    add("ПК, поступившие в отчётном году", pc["new"], "шт.", "ВПО-2, Р2_1")
    add("Ноутбуки и портативные ПК", pc["laptops"], "шт.", "ВПО-2, Р2_1")
    add("Планшетные компьютеры", pc["tablets"], "шт.", "ВПО-2, Р2_1")
    if students_vo:
        add("Студентов ВО (бакалавриат+специалитет+магистратура)", students_vo, "чел.", "ВПО-2, tab19")
        add("ПК на 100 студентов ВО", pc["total"] / students_vo * 100, "шт./100 чел.", "расчёт")
        add("ПК учебного назначения на 100 студентов ВО", pc["edu"] / students_vo * 100, "шт./100 чел.", "расчёт")
        add("Доля ПК, поступивших в отчётном году", pc["new"] / pc["total"] * 100 if pc["total"] else None, "%", "расчёт")

    for label, value in total22.items():
        add("Организаций: " + label, value, "ед.", "ВПО-2, Р2_2")
    orgs_total = max((v for k, v in total22.items() if "(доступно" not in k), default=None)
    if orgs_total:
        add("Организаций в форме (макс. по категориям Р2_2)", orgs_total, "ед.", "ВПО-2, Р2_2")

    speeds = {
        1: "ниже 256 Кбит/с", 2: "256–511 Кбит/с", 3: "512–999 Кбит/с", 4: "1,0–1,9 Мбит/с",
        5: "2,0–29,9 Мбит/с", 6: "30,0–49,9 Мбит/с", 7: "50,0–99,9 Мбит/с", 8: "выше 100 Мбит/с", 9: "не используется",
    }
    for code, label in speeds.items():
        if code in total23:
            add("Организаций с максимальной скоростью Интернета: " + label, total23[code], "ед.", "ВПО-2, Р2_3")

    with open(OUT_CSV, "w", newline="", encoding="utf-8-sig") as fh:
        w = csv.DictWriter(fh, fieldnames=["metric", "value", "unit", "source"], delimiter=";")
        w.writeheader()
        w.writerows(rows)
    print("готово:", OUT_CSV, f"({len(rows)} строк)")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
